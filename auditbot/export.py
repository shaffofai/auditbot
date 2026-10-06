"""A range of rows as the three things the bot sends: a summary message, an
Excel workbook and a JSON file.

Excel: one sheet "Xulosa" (the summary), then one sheet per source, written in
openpyxl's write-only mode so twenty thousand rows do not need gigabytes. Every
value that came from a request is written as TEXT: a path or a message starting
with "=" would otherwise become a formula the moment the file is opened.

JSON: the rows exactly as the APIs returned them, for tools rather than eyes.
"""

from __future__ import annotations

import html
import io
import json
from collections import Counter
from datetime import datetime, tzinfo
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from auditbot.sources import parse_time

EXCEL_CELL_MAX = 32_767

# (header, key, width). Keys are the API's own field names.
PLATFORM_COLUMNS = [
    ("ID", "id", 9), ("Vaqt (Toshkent)", "created_at", 20), ("Tur", "kind", 9),
    ("Hodisa", "event", 16), ("Daraja", "level", 9), ("Manba", "source", 11),
    ("Xizmat", "service", 11), ("Foydalanuvchi", "username", 26), ("Tomon", "user_side", 7),
    ("Admin", "is_admin", 7), ("Tasdiqlangan", "authenticated", 7), ("IP", "ip", 15),
    ("Usul", "method", 7), ("Yo'l", "path", 34), ("So'rov qatori", "query", 20),
    ("Holat", "status_code", 7), ("Davomiylik, ms", "duration_ms", 9), ("Xabar", "message", 50),
    ("Takror", "repeat_count", 7), ("Oxirgi marta", "last_seen_at", 20),
    ("So'rov ID", "request_id", 34), ("Sessiya", "session_id", 20), ("Fayllar", "files", 7),
    ("User-Agent", "user_agent", 30), ("Meta", "meta", 40), ("Izoh", "note", 20),
]
TENDER_COLUMNS = [
    ("ID", "id", 9), ("Vaqt (Toshkent)", "yaratildi", 20), ("Tur", "tur", 8),
    ("Hodisa", "hodisa", 16), ("Daraja", "daraja", 9), ("Manba", "manba", 8),
    ("Usul", "usul", 7), ("Yo'l", "yol", 30), ("So'rov qatori", "sorov_qatori", 20),
    ("Holat", "holat_kodi", 7), ("Davomiylik, ms", "davomiylik_ms", 9), ("IP", "ip", 15),
    ("Login", "login", 18), ("Tasdiqlangan", "tasdiqlangan", 7), ("Xabar", "xabar", 50),
    ("So'rov tanasi, belgi", "sorov_tanasi_uzunligi", 9),
    ("Javob tanasi, belgi", "javob_tanasi_uzunligi", 9), ("Qo'shimcha", "qoshimcha", 40),
    ("So'rov ID", "sorov_id", 34),
]
PLATFORM_BODIES = [("So'rov tanasi", "request_body", 50), ("Javob tanasi", "response_body", 50)]
TENDER_BODIES = [("So'rov tanasi", "sorov_tanasi", 50), ("Javob tanasi", "javob_tanasi", 50)]
TIME_KEYS = {"created_at", "last_seen_at", "yaratildi"}

PLATFORM_KINDS = [("request", "so'rov"), ("auth", "kirish"), ("client", "brauzer"),
                  ("log", "log"), ("manual", "qo'lda")]
TENDER_KINDS = [("sorov", "so'rov"), ("kirish", "kirish"), ("log", "log"), ("amal", "amal")]


def local_label(value: datetime, tz: tzinfo) -> str:
    return value.astimezone(tz).strftime("%d.%m.%Y %H:%M")


def file_stem(start: datetime, end: datetime, tz: tzinfo) -> str:
    return (f"loglar_{start.astimezone(tz).strftime('%Y%m%d-%H%M')}"
            f"_{end.astimezone(tz).strftime('%Y%m%d-%H%M')}")


# ---------- summary ----------

def _stats(rows: list[dict], *, kind_key: str, level_key: str, status_key: str,
           user_key: str, failed_event: str | None) -> dict[str, Any]:
    kinds = Counter(r.get(kind_key) for r in rows)
    statuses = [r.get(status_key) for r in rows if isinstance(r.get(status_key), int)]
    return {
        "rows": len(rows),
        "with_repeats": sum(max(int(r.get("repeat_count") or 1), 1) for r in rows),
        "kinds": kinds,
        "errors": sum(1 for r in rows if r.get(level_key) == "error"),
        "s5xx": sum(1 for s in statuses if s >= 500),
        "s4xx": sum(1 for s in statuses if 400 <= s < 500),
        "failed_logins": (sum(1 for r in rows if r.get("event") == failed_event)
                          if failed_event else kinds.get("kirish", 0)),
        "users": len({r.get(user_key) for r in rows if r.get(user_key)}),
    }


def _source_lines(title: str, stats: dict, kinds: list[tuple[str, str]], truncated: bool,
                  limit: int) -> list[str]:
    head = f"<b>{html.escape(title, quote=False)}</b>: {stats['rows']:,} ta yozuv".replace(",", " ")
    if stats["with_repeats"] > stats["rows"]:
        head += f" (takrorlar bilan {stats['with_repeats']:,})".replace(",", " ")
    lines = [head]
    if stats["rows"]:
        lines.append("  " + " · ".join(f"{label} {stats['kinds'].get(key, 0)}"
                                       for key, label in kinds))
        lines.append(f"  xatolar: {stats['errors']} · 5xx: {stats['s5xx']} · 4xx: {stats['s4xx']}"
                     f" · kirishda xato: {stats['failed_logins']}"
                     f" · foydalanuvchilar: {stats['users']}")
    if truncated:
        lines.append(f"  ⚠️ cheklovga yetdi: faqat {limit:,} ta yozuv olindi — oraliqni "
                     f"qisqartiring".replace(",", " "))
    return lines


def summary(*, start: datetime, end: datetime, tz: tzinfo, filters_label: str,
            platform: tuple[list[dict], bool] | None, tender: tuple[list[dict], bool] | None,
            limit: int, now: datetime, notes: Iterable[str] = ()) -> str:
    lines = [f"📋 <b>Loglar</b>: {local_label(start, tz)} — {local_label(end, tz)} (Toshkent)"]
    if filters_label:
        lines.append(f"Filtr: {html.escape(filters_label, quote=False)}")
    if platform is not None:
        stats = _stats(platform[0], kind_key="kind", level_key="level", status_key="status_code",
                       user_key="username", failed_event="login_failed")
        lines += _source_lines("Platforma", stats, PLATFORM_KINDS, platform[1], limit)
    if tender is not None:
        stats = _stats(tender[0], kind_key="tur", level_key="daraja", status_key="holat_kodi",
                       user_key="login", failed_event=None)
        lines += _source_lines("tender-v2", stats, TENDER_KINDS, tender[1], limit)
    lines += [html.escape(n, quote=False) for n in notes]
    if (now - end).total_seconds() < 600:
        lines.append("ℹ️ Oraliq hozirgacha davom etadi: tugamagan so'rovlar va kechikkan "
                     "brauzer xabarlari keyinroq qo'shilishi mumkin.")
    return "\n".join(lines)


# ---------- Excel ----------

def _cell_value(value: Any, key: str, tz: tzinfo) -> Any:
    if value is None:
        return None
    if key in TIME_KEYS:
        moment = parse_time(value)
        return moment.astimezone(tz).replace(tzinfo=None) if moment else str(value)
    if isinstance(value, bool):
        return "ha" if value else "yo'q"
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    if len(text) > EXCEL_CELL_MAX:
        text = text[:EXCEL_CELL_MAX - 30] + " …[Excel uchun qirqildi]"
    return text


def _text_cell(ws, value: Any) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=value)
    if isinstance(value, str):
        cell.data_type = "s"              # never a formula, whatever the text starts with
    elif isinstance(value, datetime):
        cell.number_format = "yyyy-mm-dd hh:mm:ss"
    return cell


def _sheet(wb: Workbook, title: str, columns: list[tuple[str, str, int]], rows: list[dict],
           tz: tzinfo) -> None:
    ws = wb.create_sheet(title)
    for index, (_, _, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    bold = Font(bold=True)
    header = []
    for name, _, _ in columns:
        cell = WriteOnlyCell(ws, value=name)
        cell.font = bold
        header.append(cell)
    ws.append(header)
    for row in rows:
        ws.append([_text_cell(ws, _cell_value(row.get(key), key, tz)) for _, key, _ in columns])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"


def workbook(*, summary_text: str, tz: tzinfo, platform: list[dict] | None,
             tender: list[dict] | None, bodies: bool) -> bytes:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Xulosa")
    ws.column_dimensions["A"].width = 110
    for line in html.unescape(_strip_tags(summary_text)).split("\n"):
        ws.append([_text_cell(ws, line)])
    if platform is not None:
        _sheet(wb, "Platforma", PLATFORM_COLUMNS + (PLATFORM_BODIES if bodies else []),
               platform, tz)
    if tender is not None:
        _sheet(wb, "tender-v2", TENDER_COLUMNS + (TENDER_BODIES if bodies else []), tender, tz)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _strip_tags(text: str) -> str:
    return text.replace("<b>", "").replace("</b>", "")


# ---------- JSON ----------

def _dumps(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, default=str, separators=(",", ":"))


def json_file(*, start: datetime, end: datetime, tz: tzinfo, filters_label: str,
              platform: tuple[list[dict], bool] | None,
              tender: tuple[list[dict], bool] | None) -> bytes:
    """One row per line: compact enough to stay under Telegram's 50 MB where an
    indented file would not, and still readable in an editor."""
    period = {"dan": start.astimezone(tz).isoformat(), "gacha": end.astimezone(tz).isoformat()}
    parts = ['{"oraliq":', _dumps(period), ',"filtrlar":', _dumps(filters_label or None)]
    for key, value in (("platforma", platform), ("tender_v2", tender)):
        if value is None:
            continue
        rows, truncated = value
        parts.append(f',\n"{key}":{{"yozuvlar_soni":{len(rows)},"qirqilgan":{_dumps(truncated)},'
                     '"yozuvlar":[\n')
        parts.append(",\n".join(_dumps(row) for row in rows))
        parts.append("\n]}")
    parts.append("}\n")
    return "".join(parts).encode("utf-8")
