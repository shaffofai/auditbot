"""The bot end to end against the three fakes: commands in, messages and files out;
alerts from new rows; the cursor kept across a restart."""

import asyncio
import io
import json
import logging
from datetime import timedelta

import httpx
from openpyxl import load_workbook

from auditbot.bot import Bot
from auditbot.sources import Platform, Tender
from auditbot.telegram import Telegram
from fakes import T0, FakePlatform, FakeTelegram, FakeTender

NOW = T0 + timedelta(hours=5)          # 14:00 in Tashkent; the fakes' rows start at 09:00


class Clock:
    def __init__(self):
        self.t = 5000.0

    def __call__(self):
        return self.t


class World:
    def __init__(self, settings, *, tender=True):
        self.settings = settings
        self.tg = FakeTelegram(settings.telegram_token)
        self.platform = FakePlatform(settings.platform_user, settings.platform_password)
        self.tender = FakeTender(settings.tender_user, settings.tender_password)
        self.clock = Clock()
        self.with_tender = tender
        # As an admin left them in the dashboard: a person, and the alert group.
        self.platform.chat(111, title="Ali Valiyev")
        self.platform.chat(-1002, title="Audit", alerts=True)

    def bot(self) -> Bot:
        s = self.settings
        return Bot(s,
                   Telegram(s.telegram_token, s.telegram_api,
                            httpx.AsyncClient(transport=self.tg.transport())),
                   Platform(s.platform_url, s.platform_user, s.platform_password,
                            httpx.AsyncClient(transport=self.platform.transport())),
                   Tender(s.tender_url, s.tender_user, s.tender_password,
                          httpx.AsyncClient(transport=self.tender.transport()))
                   if self.with_tender else None,
                   clock=self.clock, now=lambda: NOW)


def chat_of(chat_id):
    if chat_id > 0:
        return {"id": chat_id, "type": "private", "first_name": "Ali", "last_name": "Valiyev",
                "username": "ali_v"}
    return {"id": chat_id, "type": "supergroup", "title": "Audit"}


def message(chat_id, text, update_id=1):
    return {"update_id": update_id, "message": {"chat": chat_of(chat_id), "text": text}}


async def settle(bot):
    """Let the export tasks the bot started finish."""
    for _ in range(50):
        await asyncio.sleep(0)
        if not bot.exporting:
            return
    while bot.exporting:
        await asyncio.sleep(0.01)


def run_commands(world, *texts, chat_id=111):
    async def go():
        bot = world.bot()
        for i, text in enumerate(texts):
            await bot.handle(message(chat_id, text, i))
            await settle(bot)
        return bot
    return asyncio.run(go())


def test_export_sends_summary_xlsx_and_json(make_settings):
    world = World(make_settings())
    world.platform.add(kind="request", status_code=500, source="gasn", username="gov",
                       path="/api-v2/gasn/back/ask")
    world.platform.add(kind="auth", event="login_failed", level="warning", username="bob")
    world.platform.add(kind="client", event="console.warn", level="warning", source="web")
    world.tender.add(tur="kirish", hodisa="kirish_rad", login="tender", holat_kodi=401)

    run_commands(world, "/logs today")

    texts = world.tg.messages(111)
    assert texts[0].startswith("⏳ Tayyorlanmoqda: 05.10.2026 00:00 — 05.10.2026 14:00")
    summary = texts[1]
    assert "<b>Platforma</b>: 3 ta yozuv" in summary and "5xx: 1" in summary
    assert "<b>tender-v2</b>: 1 ta yozuv" in summary
    names = [d["filename"] for d in world.tg.documents]
    assert names == ["loglar_20261005-0000_20261005-1400.xlsx",
                     "loglar_20261005-0000_20261005-1400.json"]
    wb = load_workbook(io.BytesIO(world.tg.documents[0]["content"]))
    assert wb.sheetnames == ["Xulosa", "Platforma", "tender-v2"]
    assert wb["Platforma"].max_row == 4 and wb["tender-v2"].max_row == 2
    doc = json.loads(world.tg.documents[1]["content"])
    assert doc["platforma"]["yozuvlar_soni"] == 3 and doc["tender_v2"]["yozuvlar_soni"] == 1
    # The range reached both APIs in UTC.
    first = world.platform.calls[0]["params"]
    assert first["from"] == "2026-10-04T19:00:00+00:00" and first["to"] == "2026-10-05T09:00:00+00:00"
    assert world.tender.calls[0]["params"]["dan"] == "2026-10-04T19:00:00+00:00"


def test_export_with_filters_and_bodies(make_settings):
    world = World(make_settings(export_max_bodies=1))
    world.platform.add(level="error", request_body='{"q":1}', response_body="boom")
    world.platform.add(level="error", request_body='{"q":2}', response_body="boom2")
    world.platform.add(level="info")

    run_commands(world, "/logs today errors service=platform bodies")

    assert world.tender.calls == []
    summary = world.tg.messages(111)[1]
    assert "Filtr: errors, service=platform, bodies" in summary
    assert "Tanalar birinchi 1 ta yozuv uchun olindi." in summary
    doc = json.loads(world.tg.documents[1]["content"])
    rows = doc["platforma"]["yozuvlar"]
    assert [r["id"] for r in rows] == [1, 2]
    assert rows[0]["request_body"] == '{"q":1}' and "request_body" not in rows[1]
    assert "tender_v2" not in doc
    wb = load_workbook(io.BytesIO(world.tg.documents[0]["content"]))
    header = [c.value for c in wb["Platforma"][1]]
    assert "So'rov tanasi" in header


def test_nothing_found(make_settings):
    world = World(make_settings())
    run_commands(world, "/logs yesterday")
    assert world.tg.messages(111)[-1] == "Bu oraliqda hech narsa topilmadi."
    assert world.tg.documents == []


def test_one_source_failing_does_not_stop_the_other(make_settings):
    world = World(make_settings())
    world.tender.password = "a-different-one-of-24-chars"
    world.platform.add()
    run_commands(world, "/logs today")
    summary = world.tg.messages(111)[1]
    assert "tender-v2: o'qib bo'lmadi (login yoki parol rad etildi (401))" in summary
    assert len(world.tg.documents) == 2


def test_tender_disabled(make_settings):
    world = World(make_settings(), tender=False)
    world.platform.add()
    run_commands(world, "/logs today")
    assert "tender-v2 o'chirilgan" in world.tg.messages(111)[1]
    assert load_workbook(io.BytesIO(world.tg.documents[0]["content"])).sheetnames == ["Xulosa", "Platforma"]


def test_bad_request_and_help(make_settings):
    world = World(make_settings())
    run_commands(world, "/logs tomorrow", "/help", "/logs", "/start@ShaffofAuditBot")
    texts = world.tg.messages(111)
    assert texts[0].startswith("❓ Tushunilmadi: tomorrow")
    assert "<b>/logs</b>" in texts[1] and "birinchi 300 ta" in texts[1]
    assert texts[2].startswith("Qaysi oraliq?")
    assert "<b>/logs</b>" in texts[3]
    keyboards = [m.get("reply_markup") for m in world.tg.sent]
    assert keyboards[1]["inline_keyboard"][0][0] == {"text": "Oxirgi 1 soat", "callback_data": "q:1h"}


def test_quick_button_runs_an_export(make_settings):
    world = World(make_settings())
    world.platform.add()

    async def go():
        bot = world.bot()
        await bot.handle({"update_id": 1, "callback_query": {
            "id": "cb1", "data": "q:today", "message": {"chat": {"id": 111}}}})
        await settle(bot)
    asyncio.run(go())
    assert len(world.tg.documents) == 2


def test_a_stranger_becomes_a_request_and_waits(make_settings, caplog):
    world = World(make_settings())
    world.platform.add()
    caplog.set_level(logging.INFO)
    run_commands(world, "/logs today", "/help", "/id", "/id", chat_id=999)
    row = world.platform.chats[999]
    assert row["status"] == "pending" and row["title"] == "Ali Valiyev" and row["username"] == "ali_v"
    texts = world.tg.messages(999)
    assert texts[0].startswith("⏳ Botdan foydalanish so'rovingiz administratorlarga yuborildi")
    assert "<code>999</code>" in texts[0]
    assert texts[1:] == ["Shu chatning ID raqami: <code>999</code>"]   # reminder once; /id once per 10 s
    notices = [t for t in world.tg.messages(-1002) if "so'rovi" in t]
    assert len(notices) == 1 and "Ali Valiyev" in notices[0] and "shaxsiy chat" in notices[0]
    assert "@ali_v" in notices[0] and "<code>999</code>" in notices[0]
    assert world.tg.documents == [] and world.platform.calls == []      # no log was read
    assert "/logs from chat 999 (unknown)" in caplog.text
    # A forged button press from a stranger does nothing either.
    async def go():
        bot = world.bot()
        await bot.handle({"update_id": 9, "callback_query": {
            "id": "cb", "data": "q:today", "message": {"chat": chat_of(999)}}})
        await settle(bot)
    asyncio.run(go())
    assert world.tg.documents == []


def test_a_person_writing_without_a_command_is_a_request_too(make_settings):
    world = World(make_settings())
    run_commands(world, "salom", chat_id=998)
    assert world.platform.chats[998]["status"] == "pending"
    assert world.tg.messages(998)[0].startswith("⏳")


def test_group_chatter_is_ignored(make_settings):
    world = World(make_settings())
    run_commands(world, "hello everyone", chat_id=-777)
    assert world.tg.sent == [] and world.platform.chats.get(-777) is None


def test_approval_works_on_the_next_command(make_settings):
    world = World(make_settings())
    world.platform.add()

    async def go():
        bot = world.bot()
        await bot.handle(message(555, "/logs today", 1))
        await settle(bot)
        world.platform.chats[555]["status"] = "approved"        # the admin, in the dashboard
        world.clock.t += 6
        await bot.handle(message(555, "/logs today", 2))
        await settle(bot)
    asyncio.run(go())
    assert len([d for d in world.tg.documents if d["chat_id"] == 555]) == 2


def test_revoked_within_a_minute_and_blocked_is_silent(make_settings):
    world = World(make_settings())
    world.platform.add()

    async def go():
        bot = world.bot()
        await bot.handle(message(111, "/status", 1))
        world.platform.chats[111]["status"] = "blocked"
        world.clock.t += 61
        await bot.handle(message(111, "/logs today", 2))
        await bot.handle(message(111, "/id", 3))
        await settle(bot)
    asyncio.run(go())
    texts = world.tg.messages(111)
    assert len(texts) == 2 and texts[0].startswith("🩺") and texts[1].startswith("Shu chatning ID")
    assert world.tg.documents == []
    assert not [t for t in world.tg.messages(-1002) if "so'rovi" in t]   # no new request either


def test_bot_added_to_a_group_and_removed(make_settings):
    world = World(make_settings())
    group = {"id": -100777, "type": "supergroup", "title": "Ops <team>"}

    def member(status):
        return {"update_id": 1, "my_chat_member": {
            "chat": group, "from": {"id": 5}, "old_chat_member": {}, 
            "new_chat_member": {"status": status, "user": {"id": 42, "is_bot": True}}}}

    async def go():
        bot = world.bot()
        await bot.handle(member("member"))
        await bot.handle(member("left"))
    asyncio.run(go())
    row = world.platform.chats[-100777]
    assert row["status"] == "pending" and row["bot_member"] is False and row["title"] == "Ops <team>"
    assert world.tg.messages(-100777)[0].startswith("⏳")
    notice = [t for t in world.tg.messages(-1002) if "so'rovi" in t][0]
    assert "Ops &lt;team&gt; — guruh" in notice


def test_group_upgraded_to_supergroup_keeps_its_approval(make_settings):
    world = World(make_settings())
    world.platform.chat(-500, chat_type="group", title="Old", alerts=True)
    upgrade = {"update_id": 1, "message": {"chat": {"id": -100500, "type": "supergroup",
                                                    "title": "Old"},
                                           "migrate_from_chat_id": -500}}

    async def go():
        bot = world.bot()
        await bot.handle(upgrade)
        return bot
    bot = asyncio.run(go())
    assert -500 not in world.platform.chats
    assert world.platform.chats[-100500]["status"] == "approved"
    assert bot.chats.known[-100500]["alerts"] is True and -500 not in bot.chats.known


def test_an_alert_chat_that_blocked_the_bot_is_marked_and_skipped(make_settings):
    world = World(make_settings())
    world.platform.chat(222, alerts=True)
    world.tg.blocked_by.add(222)

    async def go():
        bot = world.bot()
        await bot.broadcast(["first"])
        await bot.broadcast(["second"])
    asyncio.run(go())
    assert world.platform.chats[222]["bot_member"] is False
    assert world.tg.messages(-1002) == ["first", "second"]
    reports = [c for c in world.platform.chat_calls if c["path"].endswith("/seen")]
    assert len(reports) == 1 and reports[0]["body"]["bot_member"] is False


def test_gateway_down_alerts_still_reach_the_last_known_chats(make_settings):
    world = World(make_settings())

    async def first():
        await world.bot().chats.refresh()             # a run that read the list, then stopped
    asyncio.run(first())
    world.platform.chats_fail_with = 500

    async def second():
        bot = world.bot()
        await bot.broadcast(["⚠️ platform down"])
        await bot.handle(message(997, "/start", 1))   # a stranger meanwhile
    asyncio.run(second())
    assert world.tg.messages(-1002) == ["⚠️ platform down"]
    assert world.tg.messages(997)[0].startswith("⚠️ Bot hozir platformaga ulana olmayapti")


def test_no_alert_chat_means_no_alert(make_settings, caplog):
    world = World(make_settings())
    world.platform.chats[-1002]["alerts"] = False
    caplog.set_level(logging.INFO)

    async def go():
        await world.bot().broadcast(["x"])
    asyncio.run(go())
    assert world.tg.sent == [] and "no approved chat has alerts on" in caplog.text


def test_one_export_per_chat_at_a_time(make_settings):
    world = World(make_settings())
    world.platform.add()

    async def go():
        bot = world.bot()
        await bot.handle(message(111, "/logs today", 1))
        await bot.handle(message(111, "/logs yesterday", 2))
        await settle(bot)
    asyncio.run(go())
    assert any("Oldingi so'rov hali tayyorlanmoqda" in t for t in world.tg.messages(111))
    assert len(world.tg.documents) == 2


def test_status(make_settings):
    world = World(make_settings())
    world.platform.add()
    run_commands(world, "/status")
    text = world.tg.messages(111)[0]
    assert "Platforma: ✅ ishlayapti, oxirgi yozuv #1" in text
    assert "tender-v2: ✅ ishlayapti, oxirgi yozuv #0" in text
    assert "Ogohlantirishlar: yoqilgan" in text
    assert "Chatlar: 2 ta ruxsat etilgan, 1 tasi ogohlantirish oladi, 0 ta so'rov kutmoqda" in text


def test_poll_loop_reads_updates_and_moves_the_offset(make_settings):
    world = World(make_settings())
    world.tg.updates = [message(111, "/id", 41), message(111, "/status", 42)]

    async def go():
        bot = world.bot()
        task = asyncio.create_task(bot.poll())
        for _ in range(100):
            await asyncio.sleep(0.01)
            if len(world.tg.messages(111)) >= 2:
                break
        task.cancel()
        return bot
    bot = asyncio.run(go())
    assert bot.offset == 43
    assert world.tg.messages(111)[0].startswith("Shu chatning ID raqami")


# ---------- alerts ----------

def test_alerts_start_at_the_end_then_report_new_rows(make_settings):
    world = World(make_settings())
    world.platform.add(status_code=500)                  # before the bot: never alerted
    world.tender.add(tur="sorov", holat_kodi=500)

    async def go():
        bot = world.bot()
        for f in bot.followers:
            await bot.follow_one(f)
        assert world.tg.sent == []
        assert bot.state.get("platform_after_id") == 1 and bot.state.get("tender_after_id") == 1

        world.platform.add(status_code=502, source="fasad", path="/api-v2/fasad/front/report",
                           method="POST", meta={"exception": "TimeoutError"})
        world.platform.add(kind="client", event="js.error", source="web",
                           message="TypeError: x is undefined")
        world.tender.add(tur="sorov", usul="POST", yol="/check", holat_kodi=500,
                         qoshimcha={"istisno": "OperationalError"})
        for f in bot.followers:
            await bot.follow_one(f)
        return bot
    bot = asyncio.run(go())

    texts = world.tg.messages(-1002)
    assert len(texts) == 3 and world.tg.messages(111) == []     # only to chats with alerts on
    assert "Server xatosi 502" in texts[0] and "TimeoutError" in texts[0]
    assert "Brauzer xatolari" in texts[1]
    assert "tender-v2 server xatosi 500" in texts[2]

    saved = json.load(open(world.settings.state_path))
    assert {k: saved[k] for k in ("platform_after_id", "tender_after_id")} == {
        "platform_after_id": 3, "tender_after_id": 2}
    assert saved["chats"]["-1002"]["alerts"] is True           # the chat list is kept too


def test_alert_cursor_survives_a_restart(make_settings):
    world = World(make_settings())
    world.platform.add()

    async def first():
        bot = world.bot()
        await bot.follow_one(bot.followers[0])
    asyncio.run(first())
    world.platform.add(status_code=500)                  # while the bot was down

    async def second():
        bot = world.bot()
        await bot.follow_one(bot.followers[0])
    asyncio.run(second())
    assert len(world.tg.messages(-1002)) == 1 and "Server xatosi 500" in world.tg.messages(-1002)[0]


def test_long_outage_skips_ahead(make_settings):
    world = World(make_settings(alert_catchup_rows=400))
    world.platform.add()

    async def go():
        bot = world.bot()
        await bot.follow_one(bot.followers[0])
        for _ in range(1000):
            world.platform.add(status_code=500, path=f"/p{_ % 3}")
        await bot.follow_one(bot.followers[0])
        return bot
    bot = asyncio.run(go())
    texts = world.tg.messages(-1002)
    assert any("tadan ko'p yozuv qo'shildi" in t for t in texts)
    assert bot.state.get("platform_after_id") == 1001
    tail_calls = [c for c in world.platform.calls if "after_id" in c["params"]]
    assert len(tail_calls) < 10


def test_source_down_and_back(make_settings):
    world = World(make_settings())
    world.platform.add()

    async def go():
        bot = world.bot()
        f = bot.followers[0]
        await bot.follow_one(f)
        world.platform.fail_with = 500
        for _ in range(3):
            await bot.follow_one(f)
        assert f.next_at > world.clock.t                 # backing off
        world.platform.fail_with = None
        await bot.follow_one(f)
    asyncio.run(go())
    texts = world.tg.messages(-1002)
    assert len(texts) == 2
    assert "Platforma</b> javob bermayapti" in texts[0] and "HTTP 500" in texts[0]
    assert "yana javob bermoqda" in texts[1]


def test_refused_credentials_wait_ten_minutes(make_settings):
    """Someone changed the bot's password on the platform: the follow loop backs
    off, and the alert still reaches the chats the bot knew before."""
    world = World(make_settings())

    async def go():
        bot = world.bot()
        await bot.chats.refresh()
        world.platform.password = "changed-in-the-dashboard"
        f = bot.followers[0]
        await bot.follow_one(f)
        return f
    f = asyncio.run(go())
    assert f.next_at - world.clock.t == 600
    assert len(world.platform.calls) == 1
    assert "bot kira olmadi" in world.tg.messages(-1002)[0]


def test_token_never_logged(make_settings, caplog):
    caplog.set_level(logging.DEBUG)
    world = World(make_settings())

    def broken(request):
        raise httpx.ConnectError("boom", request=request)

    async def go():
        bot = world.bot()
        bot.tg = Telegram(world.settings.telegram_token, world.settings.telegram_api,
                          httpx.AsyncClient(transport=httpx.MockTransport(broken)))
        task = asyncio.create_task(bot.poll())
        await asyncio.sleep(0.05)
        task.cancel()
        await bot._send_quietly(111, "x")
    asyncio.run(go())
    assert "getUpdates failed: network error (ConnectError)" in caplog.text
    assert world.settings.telegram_token not in caplog.text


def test_user_text_is_escaped_in_replies(make_settings):
    world = World(make_settings())
    run_commands(world, "/logs <b>")
    assert world.tg.messages(111) == ["❓ Tushunilmadi: &lt;b&gt;\n\nYordam: /help"]


def test_markup_telegram_refuses_goes_again_as_plain_text(make_settings):
    world = World(make_settings())

    async def go():
        await world.bot().tg.send(111, "<b>qalin</b> &amp; <x>")
    asyncio.run(go())
    assert world.tg.sent[-1]["text"] == "qalin & " and "parse_mode" not in world.tg.sent[-1]


def test_big_json_is_zipped_and_too_big_files_explained(make_settings, monkeypatch):
    import auditbot.bot as botmodule
    monkeypatch.setattr(botmodule, "FILE_LIMIT", 6000)
    world = World(make_settings())
    for i in range(200):
        world.platform.add(message="the same words again " * 3, username=f"user{i % 5}@misol.uz")
    run_commands(world, "/logs today service=platform")
    names = [d["filename"] for d in world.tg.documents]
    assert names == ["loglar_20261005-0000_20261005-1400.json.zip"]
    import zipfile
    with zipfile.ZipFile(io.BytesIO(world.tg.documents[0]["content"])) as zf:
        doc = json.loads(zf.read("loglar_20261005-0000_20261005-1400.json"))
    assert doc["platforma"]["yozuvlar_soni"] == 200
    assert any("juda katta" in t and ".xlsx" in t for t in world.tg.messages(111))


def test_commands_are_registered_in_english(make_settings):
    world = World(make_settings())
    calls = []

    async def record(commands):
        calls.append([c for c, _ in commands])
        raise asyncio.CancelledError

    async def go():
        bot = world.bot()
        bot.tg.commands = record
        try:
            await bot.run()
        except asyncio.CancelledError:
            pass
    asyncio.run(go())
    assert calls == [["logs", "status", "help", "id"]]


def test_old_uzbek_commands_do_nothing(make_settings):
    world = World(make_settings())
    world.platform.add()
    run_commands(world, "/loglar bugun", "/holat", "/yordam")
    assert world.tg.sent == [] and world.tg.documents == []


def test_a_just_approved_alert_group_hears_of_the_next_request(make_settings):
    world = World(make_settings())
    world.platform.chats[-1002]["status"] = "pending"

    async def go():
        bot = world.bot()
        await bot.chats.refresh()
        world.platform.chats[-1002]["status"] = "approved"      # the admin, a moment later
        world.clock.t += 6
        await bot.handle(message(996, "/start", 1))
    asyncio.run(go())
    assert any("so'rovi" in t for t in world.tg.messages(-1002))
