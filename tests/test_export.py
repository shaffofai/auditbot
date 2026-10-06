import io
import json
from datetime import datetime, timezone

from openpyxl import load_workbook

from auditbot import export
from conftest import TZ

START = datetime(2026, 10, 5, 4, 0, tzinfo=timezone.utc)
END = datetime(2026, 10, 5, 13, 0, tzinfo=timezone.utc)
NOW = datetime(2026, 10, 6, 0, 0, tzinfo=timezone.utc)

PLATFORM = [
    {"id": 1, "created_at": "2026-10-05T04:00:01.5Z", "kind": "request", "event": "http",
     "level": "info", "source": "fasad", "username": "ali@misol.uz", "method": "POST",
     "path": "=HYPERLINK(\"http://evil\")", "status_code": 200, "is_admin": False,
     "repeat_count": 1, "meta": {"a": 1}, "message": "bad\x00\x07char"},
    {"id": 2, "created_at": "2026-10-05T04:05:00Z", "kind": "auth", "event": "login_failed",
     "level": "warning", "source": "gateway", "username": "bob", "repeat_count": 3},
    {"id": 3, "created_at": "2026-10-05T04:06:00Z", "kind": "request", "event": "http",
     "level": "error", "source": "gasn", "status_code": 502, "username": "gov",
     "repeat_count": 1, "message": "x" * 40000},
]
TENDER = [
    {"id": 7, "yaratildi": "2026-10-05T04:10:00+00:00", "tur": "kirish", "hodisa": "kirish_rad",
     "daraja": "warning", "login": "tender", "holat_kodi": 401, "qoshimcha": {"sabab": "parol"}},
]


def test_summary_counts_and_escapes():
    text = export.summary(start=START, end=END, tz=TZ, filters_label="foydalanuvchi=<a>",
                          platform=(PLATFORM, False), tender=(TENDER, True), limit=20000, now=NOW)
    assert "05.10.2026 09:00 — 05.10.2026 18:00" in text
    assert "&lt;a&gt;" in text
    assert "<b>Platforma</b>: 3 ta yozuv (takrorlar bilan 5)" in text
    assert "xatolar: 1 · 5xx: 1 · 4xx: 0 · kirishda xato: 1 · foydalanuvchilar: 3" in text
    assert "<b>tender-v2</b>: 1 ta yozuv" in text
    assert "kirishda xato: 1" in text.split("tender-v2")[1]
    assert "cheklovga yetdi" in text.split("tender-v2")[1]
    assert "hozirgacha" not in text


def test_summary_warns_about_an_open_range():
    text = export.summary(start=START, end=NOW, tz=TZ, filters_label="", platform=([], False),
                          tender=None, limit=20000, now=NOW)
    assert "hozirgacha" in text and "tender-v2" not in text


def test_workbook_sheets_text_and_times():
    data = export.workbook(summary_text="<b>Loglar</b> &amp; xulosa", tz=TZ, platform=PLATFORM,
                           tender=TENDER, bodies=False)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Xulosa", "Platforma", "tender-v2"]
    assert wb["Xulosa"]["A1"].value == "Loglar & xulosa"
    ws = wb["Platforma"]
    header = [c.value for c in ws[1]]
    assert header[:3] == ["ID", "Vaqt (Toshkent)", "Tur"]
    assert "So'rov tanasi" not in header
    row = {h: c for h, c in zip(header, ws[2])}
    assert row["Yo'l"].value == '=HYPERLINK("http://evil")'
    assert row["Yo'l"].data_type == "s"                       # text, never a formula
    assert row["Vaqt (Toshkent)"].value == datetime(2026, 10, 5, 9, 0, 1, 500000)
    assert row["Xabar"].value == "badchar"
    assert row["Admin"].value == "yo'q"
    assert row["Meta"].value == '{"a": 1}'
    long_row = {h: c.value for h, c in zip(header, ws[4])}
    assert len(long_row["Xabar"]) <= export.EXCEL_CELL_MAX
    assert ws.auto_filter.ref == f"A1:{ws.cell(1, len(header)).column_letter}4"
    t = wb["tender-v2"]
    t_header = [c.value for c in t[1]]
    t_row = {h: c.value for h, c in zip(t_header, t[2])}
    assert t_row["Vaqt (Toshkent)"] == datetime(2026, 10, 5, 9, 10)
    assert t_row["Qo'shimcha"] == '{"sabab": "parol"}'


def test_workbook_with_bodies_has_body_columns():
    rows = [dict(PLATFORM[0], request_body='{"x":1}', response_body="ok")]
    wb = load_workbook(io.BytesIO(export.workbook(summary_text="s", tz=TZ, platform=rows,
                                                  tender=None, bodies=True)))
    header = [c.value for c in wb["Platforma"][1]]
    assert header[-2:] == ["So'rov tanasi", "Javob tanasi"]
    assert wb.sheetnames == ["Xulosa", "Platforma"]


def test_json_file_keeps_rows_as_returned():
    doc = json.loads(export.json_file(start=START, end=END, tz=TZ, filters_label="",
                                      platform=(PLATFORM, False), tender=(TENDER, True)))
    assert doc["oraliq"] == {"dan": "2026-10-05T09:00:00+05:00", "gacha": "2026-10-05T18:00:00+05:00"}
    assert doc["platforma"]["yozuvlar_soni"] == 3 and doc["platforma"]["yozuvlar"][0] == PLATFORM[0]
    assert doc["tender_v2"]["qirqilgan"] is True
    assert doc["filtrlar"] is None


def test_file_stem():
    assert export.file_stem(START, END, TZ) == "loglar_20261005-0900_20261005-1800"
